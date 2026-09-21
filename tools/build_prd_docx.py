from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Reins_PRD_v1_1.md"
OUTPUT = ROOT / "output" / "docx" / "Reins_PRD_v1_1.docx"
TEMPLATE_SOURCE: Path | None = None
VERSION = "1.1"
SUBTITLE = "Independent Policy Enforcement and Authorization Evidence for AI Agent Spend"
DOCUMENT_CONTROL_LABEL = "Version"
DOCUMENT_CONTROL_VALUE = VERSION
FOOTER_DOCUMENT_LABEL = f"v{VERSION}"
STATUS = "Evidence based revision for validation"
DOCUMENT_DATE = "16 September 2026"
TARGET_LAUNCH = "Not set"
DECISION_REQUEST = (
    "Approve the revised wedge, enforcement model, validation gates, and sandbox MVP. "
    "Production money movement and broad all rails claims remain gated on partner, security, PCI, and legal review."
)

INK = "14213D"
NAVY = "17375E"
TEAL = NAVY
PALE = "E8EEF6"
ALT = "F3F6FA"
GRID = "C7D1DF"
MUTED = "52647A"
WHITE = "FFFFFF"


def ascii_text(text: str) -> str:
    replacements = {
        "\u2014": " - ",
        "\u2013": " - ",
        "\u2011": "-",
        "\u2212": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2026": "...",
        "\u00d7": "x",
        "\u2192": "->",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def heading_text(text: str) -> str:
    text = ascii_text(text)
    text = text.replace("/", " and ").replace("&", " and ")
    text = text.replace(":", " ").replace("-", " ")
    text = re.sub(r"[^A-Za-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def set_font(run, name: str, size: float, bold: bool | None = None, color: str = INK, italic: bool = False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def mark_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_table_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_bookmark(paragraph, name: str, bookmark_id: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph_properties = paragraph._p.pPr
    insert_at = 1 if paragraph_properties is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_internal_link(paragraph, text: str, anchor: str, *, bold=False, color=TEAL, size=9.5):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Aptos")
    r_fonts.set(qn("w:hAnsi"), "Aptos")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(int(size * 2)))
    r_pr.extend([r_fonts, color_node, underline, size_node])
    if bold:
        r_pr.append(OxmlElement("w:b"))
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = ascii_text(text)
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_external_link(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Aptos")
    r_fonts.set(qn("w:hAnsi"), "Aptos")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color_node, underline])
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = ascii_text(text)
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline(paragraph, text: str, *, base_size=10.5, base_color=INK, bold_all=False):
    text = ascii_text(text)
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, "Aptos", base_size, bold=bold_all, color=base_color)
        token = match.group(0)
        if token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_external_link(paragraph, label, url)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, "Consolas", max(base_size - 0.6, 8), bold=False, color=INK)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "E9EEF5")
            run._element.get_or_add_rPr().append(shd)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, "Aptos", base_size, bold=True, color=base_color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, "Aptos", base_size, bold=bold_all, color=base_color, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, "Aptos", base_size, bold=bold_all, color=base_color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, "Aptos", 8, color=MUTED)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.font.bold = False
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    pf.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(31)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(7)
    title_p_pr = title._element.get_or_add_pPr()
    title_borders = title_p_pr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_p_pr.remove(title_borders)

    for name, size, before, after in (
        ("Heading 1", 16, 15, 6),
        ("Heading 2", 12.5, 11, 4),
        ("Heading 3", 10.8, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_after = Pt(5)


def configure_sections(doc: Document):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)
        section.header_distance = Inches(0.32)
        section.footer_distance = Inches(0.3)


def add_header_footer(section, first_page=False):
    section.different_first_page_header_footer = first_page
    for header in (section.header, section.even_page_header):
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run("REINS  |  PRODUCT REQUIREMENTS DOCUMENT")
        set_font(run, "Aptos", 7.5, bold=True, color=MUTED)
    for footer in (section.footer, section.even_page_footer):
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Reins PRD {FOOTER_DOCUMENT_LABEL}  |  ")
        set_font(run, "Aptos", 8, color=MUTED)
        add_page_number(p)


def clear_story(story):
    root = story._element
    for child in list(root):
        root.remove(child)
    story.add_paragraph()


def clear_document_body(doc: Document):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_header_footer_stories(doc: Document):
    seen = set()
    for section in doc.sections:
        for story in (
            section.header,
            section.even_page_header,
            section.first_page_header,
            section.footer,
            section.even_page_footer,
            section.first_page_footer,
        ):
            part_name = str(story.part.partname)
            if part_name in seen:
                continue
            seen.add(part_name)
            clear_story(story)


def add_cover(doc: Document):
    for _ in range(4):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    set_font(run, "Aptos", 10, bold=True, color=TEAL)
    kicker.paragraph_format.space_after = Pt(14)

    doc.add_paragraph("Reins", style="Title")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run(SUBTITLE)
    set_font(run, "Aptos Display", 14, color=INK)

    metadata = [
        (DOCUMENT_CONTROL_LABEL, DOCUMENT_CONTROL_VALUE),
        ("Status", STATUS),
        ("Date", DOCUMENT_DATE),
        ("Contributors", "Sai Prathap Reddy Cheluri and Lekhashree Srinath Reddy"),
        ("Ownership", "Shared product requirements document"),
        ("Target launch", TARGET_LAUNCH),
        ("Audience", "Product, Engineering, Security, Payments Risk, Legal and Compliance, Partnerships"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.35)
    set_table_borders(table)
    for i, (label, value) in enumerate(metadata):
        left, right = table.rows[i].cells
        set_cell_shading(left, PALE)
        for cell in (left, right):
            set_cell_margins(cell, 100, 110, 100, 110)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = left.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, label, base_size=9.2, bold_all=False)
        p = right.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, value, base_size=9.2)

    doc.add_paragraph()
    decision = doc.add_paragraph()
    decision.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    decision.paragraph_format.left_indent = Inches(0.38)
    decision.paragraph_format.right_indent = Inches(0.38)
    decision.paragraph_format.space_before = Pt(10)
    add_inline(decision, f"Decision requested: {DECISION_REQUEST}", base_size=10.2)
    doc.add_page_break()


def extract_headings(lines):
    headings = []
    used = set()
    body_start = next(
        (index for index, line in enumerate(lines) if re.match(r"^##\s+1(?:\.|\s)", line)),
        0,
    )
    for line in lines[body_start:]:
        match = re.match(r"^(#{2,4})\s+(.+)$", line)
        if not match:
            continue
        level = len(match.group(1)) - 1
        raw = match.group(2).strip()
        display = heading_text(raw)
        slug = re.sub(r"[^A-Za-z0-9]+", "_", display).strip("_")[:34] or "section"
        base = slug
        suffix = 2
        while slug in used:
            slug = f"{base}_{suffix}"
            suffix += 1
        used.add(slug)
        headings.append({"raw": raw, "display": display, "level": level, "bookmark": slug})
    return headings


def add_contents(doc: Document, headings):
    title = doc.add_paragraph("Contents", style="Heading 1")
    add_bookmark(title, "Contents", 1)
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = intro.add_run("Select a section title to navigate within the document.")
    set_font(run, "Aptos", 9.5, color=MUTED)
    intro.paragraph_format.space_after = Pt(9)
    for item in headings:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.0 if item["level"] == 1 else 0.24 if item["level"] == 2 else 0.48)
        p.paragraph_format.space_after = Pt(1.5)
        add_internal_link(
            p,
            item["display"],
            item["bookmark"],
            bold=False,
            size=9.4 if item["level"] == 1 else 8.8,
        )
    doc.add_page_break()


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        return [rows[0]] + rows[2:], i
    return None, start


def add_table(doc: Document, rows):
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    mark_repeat_table_header(table.rows[0])

    for r_index, source_row in enumerate(rows):
        prevent_table_row_split(table.rows[r_index])
        for c_index in range(cols):
            cell = table.cell(r_index, c_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if r_index == 0:
                set_cell_shading(cell, NAVY)
            elif r_index % 2 == 0:
                set_cell_shading(cell, ALT)
            text = source_row[c_index] if c_index < len(source_row) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 18 and c_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            add_inline(
                p,
                text,
                base_size=8.2,
                base_color=WHITE if r_index == 0 else INK,
                bold_all=r_index == 0,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, code_lines):
    p = doc.add_paragraph(style="Code Block")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EEF2F7")
    p_pr.append(shd)
    run = p.add_run("\n".join(ascii_text(line) for line in code_lines))
    set_font(run, "Consolas", 7.8, color=INK)


def add_body(doc: Document, lines, headings):
    first_body_index = next(i for i, line in enumerate(lines) if re.match(r"^##\s+1(?:\.|\s)", line))
    lookup = {item["raw"]: item for item in headings}
    bookmark_id = 10
    in_references = False
    i = first_body_index
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue
        if stripped.startswith("|"):
            rows, new_i = parse_table(lines, i)
            if rows:
                add_table(doc, rows)
                i = new_i
                continue
        heading_match = re.match(r"^(#{2,4})\s+(.+)$", raw)
        if heading_match:
            source = heading_match.group(2).strip()
            item = lookup[source]
            p = doc.add_paragraph(item["display"], style=f"Heading {min(item['level'], 3)}")
            add_bookmark(p, item["bookmark"], bookmark_id)
            bookmark_id += 1
            in_references = item["display"] == "14 Source references"
            i += 1
            continue
        quote_match = re.match(r"^>\s*(.+)$", stripped)
        if quote_match:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(7)
            add_inline(p, quote_match.group(1), base_size=10.4, base_color=INK)
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        list_match = re.match(r"^(\s*)([-*]|\d+\.)\s+(.+)$", raw)
        if list_match:
            indent = min(len(list_match.group(1)) // 2, 3)
            marker = list_match.group(2)
            text = list_match.group(3)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.22 + 0.22 * indent)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(1.5 if in_references else 3)
            prefix = f"{marker} " if marker.endswith(".") else "• "
            run = p.add_run(prefix)
            set_font(run, "Aptos", 9.2 if in_references else 10, bold=False, color=INK)
            add_inline(p, text, base_size=9.2 if in_references else 10.2)
            i += 1
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, stripped, base_size=10.5)
        i += 1


def set_document_properties(doc: Document):
    props = doc.core_properties
    props.title = "Reins Product Requirements Document"
    props.subject = "Independent policy enforcement and authorization evidence for AI agent spend"
    props.author = "Sai Prathap Reddy Cheluri and Lekhashree Srinath Reddy"
    props.last_modified_by = "Sai Prathap Reddy Cheluri and Lekhashree Srinath Reddy"
    props.keywords = "agentic payments, policy enforcement, reconciliation, AP2, UCP"


def build():
    markdown = ascii_text(SOURCE.read_text(encoding="utf-8"))
    if "\u2014" in markdown:
        raise ValueError("Em dash remains in source after sanitization")
    lines = markdown.splitlines()
    headings = extract_headings(lines)

    if TEMPLATE_SOURCE is not None:
        doc = Document(TEMPLATE_SOURCE)
        clear_document_body(doc)
        clear_header_footer_stories(doc)
    else:
        doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    set_document_properties(doc)
    doc.settings.odd_and_even_pages_header_footer = True
    add_header_footer(doc.sections[0], first_page=True)
    add_cover(doc)
    add_contents(doc, headings)
    add_body(doc, lines, headings)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
